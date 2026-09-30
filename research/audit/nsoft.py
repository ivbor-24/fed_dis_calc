"""Фирменные элементы шаблона Нормасофт для сборки презентаций через python-pptx."""
import copy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

BLUE = RGBColor(0x03, 0x86, 0xFC)
INK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0xF1, 0xF3, 0xF6)
MUTED = RGBColor(0x6B, 0x72, 0x80)
PALE = RGBColor(0xE3, 0xEF, 0xFE)
LINE = RGBColor(0xC9, 0xD3, 0xE0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
F_HEAD = "VK Sans Display DemiBold"
F_BODY = "Manrope Medium"
F_NAV = "Roboto Light"


def rm(shape):
    shape._element.getparent().remove(shape._element)


def by_name(slide, prefix):
    return [s for s in slide.shapes if s.name.startswith(prefix)]


def style_run(run, size, color=INK, font=F_BODY, bold=False):
    f = run.font
    f.size = Pt(size)
    f.name = font
    f.bold = bold
    f.color.rgb = color


def text(slide, x, y, w, h, paras, size=13, color=INK, font=F_BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullet=False, space=4):
    """paras: str | list[str | list[(text, opts)]]"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space)
        runs = p if isinstance(p, list) else [(p, {})]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            style_run(r, o.get("size", size), o.get("color", color), o.get("font", font), o.get("bold", False))
        if bullet:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.24))))
            pPr.set("indent", str(-Emu(Inches(0.24))))
            buClr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(buClr, qn("a:srgbClr")).set("val", "0386FC")
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", F_BODY)
            etree.SubElement(pPr, qn("a:buChar")).set("char", "→")
    return tb


def box(slide, x, y, w, h, fill=GREY, line=None, shape=MSO_SHAPE.SNIP_1_RECTANGLE, dash=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25 if dash else 1)
        if dash:
            s.line.dash_style = 4   # MSO_LINE.DASH
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def label(shape, txt, size=12, color=WHITE, font=F_BODY, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = txt
    style_run(r, size, color, font)
    return shape


def pill(slide, x, y, w, txt, active=True, h=0.42, size=13):
    s = box(slide, x, y, w, h, fill=BLUE if active else None, line=None if active else BLUE)
    return label(s, txt, size, WHITE if active else BLUE)


def num_badge(slide, x, y, txt, size=0.5, font_size=16):
    s = box(slide, x, y, size, size, fill=BLUE)
    s.text_frame.margin_left = s.text_frame.margin_right = 0
    return label(s, txt, font_size, WHITE, F_HEAD)


def set_title(slide, title, size=24):
    t = by_name(slide, "Заголовок 1")[0]
    t.left, t.top, t.width, t.height = Inches(0.65), Inches(1.42), Inches(11.4), Inches(0.62)
    run = t.text_frame.paragraphs[0].runs[0]
    run.text = title
    run.font.size = Pt(size)
    for extra in t.text_frame.paragraphs[0].runs[1:]:
        extra._r.getparent().remove(extra._r)
    for body in by_name(slide, "Объект 2"):
        rm(body)


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def source(slide, txt, y=6.62):
    text(slide, 0.65, y, 10.9, 0.3, txt, size=9, color=MUTED)


def set_paras(shape, lines):
    """Заменить текст, сохранив оформление первого абзаца."""
    tf = shape.text_frame
    first = tf.paragraphs[0]._p
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in first.findall(qn("a:r"))[1:]:
        first.remove(r)
    first.findall(qn("a:r"))[0].find(qn("a:t")).text = lines[0]
    prev = first
    for line in lines[1:]:
        np_ = copy.deepcopy(first)
        np_.findall(qn("a:r"))[0].find(qn("a:t")).text = line
        prev.addnext(np_)
        prev = np_


def divider(slide, num, title, sub, size=34):
    big = by_name(slide, "Google Shape;529")
    numshape = min(big, key=lambda s: s.width)
    titleshape = max(big, key=lambda s: s.width)
    set_paras(numshape, [num])
    titleshape.left, titleshape.width = Inches(4.2), Inches(8.22)
    set_paras(titleshape, title)
    for p in titleshape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    subshape = by_name(slide, "TextBox 9")[0]
    subshape.left, subshape.width = Inches(5.1), Inches(7.32)
    set_paras(subshape, [sub])


def rebuild_nav(slides, n, nav, active):
    """Меню в синей плашке: nav = [(подпись, номер слайда)], active = подпись текущего раздела."""
    slide = slides[n - 1]
    for s in list(slide.shapes):
        if s.name == "nav" or s.name.startswith(("Google Shape;229", "Google Shape;230",
                                                 "Google Shape;231", "Google Shape;232")):
            rm(s)
    for s in by_name(slide, "Google Shape;269"):   # иконка «домой» -> титул
        cNvPr = s._element.find(qn("p:nvSpPr")).find(qn("p:cNvPr"))
        for h in cNvPr.findall(qn("a:hlinkClick")):
            cNvPr.remove(h)
        s.click_action.target_slide = slides[0]
    x = 1.02
    for txt, target in nav:
        w = 0.36 + 0.105 * len(txt)
        s = slide.shapes.add_shape(MSO_SHAPE.SNIP_1_RECTANGLE, Inches(x), Inches(0.47),
                                   Inches(w), Inches(0.37))
        s.name = "nav"
        s.shadow.inherit = False
        s.line.fill.background()
        on = txt == active
        if on:
            s.fill.solid()
            s.fill.fore_color.rgb = WHITE
        else:
            s.fill.background()
        s.text_frame.margin_left = s.text_frame.margin_right = 0
        label(s, txt, 11, BLUE if on else WHITE, F_NAV)
        s.click_action.target_slide = slides[target - 1]
        x += w + 0.1


def fix_img_layers(slide):
    """В шаблоне эффект яркости картинки (a14:imgLayer) ссылается на связь со слайдом, а не на hdphoto."""
    for layer in list(slide._element.iter("{http://schemas.microsoft.com/office/drawing/2010/main}imgLayer")):
        rId = layer.get(qn("r:embed"))
        rel = slide.part.rels.get(rId) if rId else None
        if rel is None or not rel.reltype.endswith("/hdphoto"):
            ext = layer.getparent().getparent()
            ext.getparent().remove(ext)


def drop_dead_slide_rels(slide):
    xml = etree.tostring(slide._element).decode()
    for rId, rel in list(slide.part.rels.items()):
        if rel.reltype.endswith("/slide") and f'"{rId}"' not in xml:
            slide.part.drop_rel(rId)


def table(slide, x, y, w, rows, col_w, header=True, size=10.5, row_h=0.36, first_bold=False):
    n = len(rows)
    tbl = slide.shapes.add_table(n, len(col_w), Inches(x), Inches(y), Inches(w), Inches(row_h * n)).table
    for i, c in enumerate(col_w):
        tbl.columns[i].width = Inches(c)
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            head = header and r == 0
            cell.fill.fore_color.rgb = BLUE if head else (GREY if r % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = v
            bold_col = first_bold and c == 0 and not head
            style_run(run, size, WHITE if head else (BLUE if bold_col else INK), F_HEAD if bold_col else F_BODY)
    return tbl
